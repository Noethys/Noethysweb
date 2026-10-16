# -*- coding: utf-8 -*-
#  Copyright (c) 2019-2026 Ivan LUCAS.
#  Noethysweb, application de gestion multi-activités.
#  Distribué sous licence GNU GPL.

"""
Importation d'événements depuis un fichier XLSX dans le paramétrage d'une activité.

Principe :
    1. Lecture du fichier : les valeurs des cellules sont converties en texte et mémorisées en session.
    2. Association des colonnes du fichier aux champs de l'événement (proposée automatiquement d'après les titres).
    3. Analyse : chaque ligne est contrôlée et transformée en un ou plusieurs événements (un par groupe).
       Les doublons (même date, nom, unité et groupe) sont ignorés, les lignes en erreur sont écartées.
    4. Importation : création des événements, des ouvertures manquantes dans le calendrier
       et recopie de la tarification d'un événement existant si demandé.
"""

import datetime, decimal, io, logging, re, unicodedata, uuid
logger = logging.getLogger(__name__)
from django.db import transaction
from core.models import Activite, Unite, Groupe, CategorieEvenement, Evenement, Ouverture, Tarif, TarifLigne
from core.utils import utils_dates, utils_historique, utils_preferences

MAX_LIGNES = 2000
MAX_TAILLE_FICHIER = 10 * 1024 * 1024    # 10 Mo
MAX_FICHIERS_SESSION = 3                 # Fichiers conservés en session (un par onglet ouvert)
MAX_CAPACITE = 100000
MAX_MONTANT = decimal.Decimal("99999999.99")   # Evenement.montant : DecimalField(max_digits=10, decimal_places=2)
MAX_ID = 2147483647
CLE_SESSION = "import_evenements"


class Erreur(Exception):
    pass


# Champs de l'événement pouvant être importés.
# "synonymes" : titres de colonnes reconnus automatiquement (comparés sans accents ni majuscules)
CHAMPS = [
    {"code": "date", "label": "Date", "obligatoire": True,
     "synonymes": ["date", "jour", "date evenement", "date de l evenement", "le"],
     "aide": "Date de l'événement. Une cellule au format date d'Excel ou un texte jj/mm/aaaa (exemple : 12/02/2027)."},
    {"code": "nom", "label": "Nom", "obligatoire": True,
     "synonymes": ["nom", "intitule", "titre", "evenement", "libelle", "nom evenement", "nom de l evenement"],
     "aide": "Nom de l'événement (200 caractères maximum). Exemple : Patinoire."},
    {"code": "description", "label": "Description",
     "synonymes": ["description", "details", "detail", "remarque", "remarques", "commentaire", "commentaires", "observations", "infos", "informations"],
     "aide": "Texte visible par les familles sur le portail (400 caractères maximum). Exemple : RDV à 13h30 devant le centre."},
    {"code": "unite", "label": "Unité",
     "synonymes": ["unite", "unites", "unite de consommation", "unite conso"],
     "aide": "Nom ou abrégé d'une unité de consommation de type Evénementiel de l'activité. Si la colonne est absente ou la cellule vide, l'unité par défaut choisie à l'étape 1 est utilisée."},
    {"code": "groupe", "label": "Groupe(s)",
     "synonymes": ["groupe", "groupes", "tranche", "tranche d age", "public"],
     "aide": "Nom ou abrégé d'un groupe de l'activité. Plusieurs groupes peuvent être séparés par des points-virgules (un événement est alors créé par groupe). "
             "Saisissez « Tous » pour créer l'événement dans tous les groupes. Si la cellule est vide, les groupes par défaut choisis à l'étape 1 sont utilisés."},
    {"code": "categorie", "label": "Catégorie",
     "synonymes": ["categorie", "categorie evenement", "categorie d evenement", "type", "type evenement"],
     "aide": "Nom d'une catégorie d'événements de l'activité (facultatif)."},
    {"code": "heure_debut", "label": "Heure de début",
     "synonymes": ["heure debut", "heure de debut", "debut", "depart", "horaire debut", "heure depart", "de"],
     "aide": "Heure au format 13:30, 13h30 ou 13h."},
    {"code": "heure_fin", "label": "Heure de fin",
     "synonymes": ["heure fin", "heure de fin", "fin", "retour", "horaire fin", "heure retour", "a"],
     "aide": "Heure au format 17:00, 17h00 ou 17h. Doit être postérieure à l'heure de début."},
    {"code": "equiv_heures", "label": "Équivalence en heures",
     "synonymes": ["equivalence heures", "equivalence en heures", "equiv heures", "equivalence", "duree"],
     "aide": "Durée au format HH:MM, utilisée par l'état global et l'état nominatif. Exemple : 03:30."},
    {"code": "capacite_max", "label": "Capacité max.",
     "synonymes": ["capacite", "capacite max", "capacite maximale", "places", "nombre de places", "nb places", "nbre places", "places max"],
     "aide": "Nombre de places maximal pour cet événement. Laissez vide ou saisissez 0 pour ne pas limiter."},
    {"code": "montant", "label": "Montant",
     "synonymes": ["montant", "prix", "tarif", "cout", "montant tarif", "tarif simple"],
     "aide": "Tarif simple de l'événement. Exemples : 6,50 ou 6.50. Laissez vide pour un événement gratuit (ou si vous utilisez la colonne Tarification à recopier)."},
    {"code": "tarif_existant", "label": "Tarification à recopier (ID)",
     "synonymes": ["id tarif", "tarif a recopier", "tarification a recopier", "id evenement tarif", "copie tarif", "recopier tarif", "id tarification",
                   "id evenement modele", "evenement modele", "modele tarif", "id modele"],
     "aide": "ID d'un événement existant de cette activité dont la tarification doit être recopiée (l'ID figure dans la colonne ID de la liste des événements). "
             "Si cet événement a un tarif simple, son montant est repris ; s'il a des tarifs avancés, ils sont dupliqués. Cette colonne est prioritaire sur la colonne Montant."},
    {"code": "visible_portail", "label": "Visible sur le portail",
     "synonymes": ["visible portail", "visible sur le portail", "portail", "visible"],
     "aide": "Oui ou Non. Si la colonne est absente ou la cellule vide, l'événement est visible sur le portail."},
]
DICT_CHAMPS = {champ["code"]: champ for champ in CHAMPS}


# ------------------------------------------- Outils de conversion -------------------------------------------

def Normaliser(texte):
    """ Minuscules, sans accents, caractères non alphanumériques remplacés par des espaces """
    texte = unicodedata.normalize("NFD", str(texte or "")).encode("ascii", "ignore").decode("ascii").lower()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", texte).split())


def Convertir_cellule(valeur):
    """ Convertit la valeur d'une cellule openpyxl en texte (les dates et heures sont normalisées) """
    if valeur is None:
        return ""
    if isinstance(valeur, datetime.datetime):
        if valeur.time() == datetime.time(0, 0):
            return valeur.date().isoformat()
        if valeur.date() in (datetime.date(1899, 12, 30), datetime.date(1899, 12, 31)):
            return valeur.strftime("%H:%M")
        return valeur.strftime("%Y-%m-%d %H:%M")
    if isinstance(valeur, datetime.date):
        return valeur.isoformat()
    if isinstance(valeur, datetime.time):
        return valeur.strftime("%H:%M")
    if isinstance(valeur, datetime.timedelta):
        minutes = int(round(valeur.total_seconds() / 60))
        return "%02d:%02d" % (minutes // 60, minutes % 60)
    if isinstance(valeur, bool):
        return "oui" if valeur else "non"
    if isinstance(valeur, float):
        return str(int(valeur)) if valeur.is_integer() else repr(valeur)
    return str(valeur).strip()


def Lire_date(texte):
    texte = texte.strip()
    if not texte:
        return None
    try:
        return _Lire_date(texte)
    except ValueError:
        raise ValueError("La date '%s' n'est pas valide (format attendu : jj/mm/aaaa)" % texte)


def _Lire_date(texte):
    # Format ISO (dates converties depuis Excel)
    resultat = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T].*)?$", texte)
    if resultat:
        annee, mois, jour = [int(x) for x in resultat.groups()]
        return datetime.date(annee, mois, jour)
    # Format français jj/mm/aaaa, jj-mm-aaaa, jj.mm.aaaa ou jj/mm/aa
    resultat = re.match(r"^(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2}|\d{4})(?:[ T]+\d{1,2}[:h]\d{2}(?::\d{2})?)?$", texte)
    if resultat:
        jour, mois, annee = [int(x) for x in resultat.groups()]
        if annee < 100:
            annee += 2000
        return datetime.date(annee, mois, jour)
    # Numéro de série Excel (cellule numérique non formatée en date)
    if re.match(r"^\d{5}(\.0+)?$", texte):
        numero = int(float(texte))
        if 20000 <= numero <= 80000:
            return datetime.date(1899, 12, 30) + datetime.timedelta(days=numero)
    raise ValueError("La date '%s' n'est pas valide (format attendu : jj/mm/aaaa)" % texte)


def Lire_heure(texte, nom="L'heure"):
    texte = texte.strip().lower().replace(" ", "")
    if not texte:
        return None
    # Datetime complet converti depuis Excel : on garde l'heure
    resultat = re.match(r"^\d{4}-\d{1,2}-\d{1,2}(\d{1,2}):(\d{2})$", texte)
    if resultat:
        return datetime.time(int(resultat.group(1)), int(resultat.group(2)))
    # 13:30, 13h30, 13h, 13:30:00, 13
    resultat = re.match(r"^(\d{1,2})(?:[:h](\d{2})?(?::\d{2})?)?$", texte)
    if resultat:
        heures, minutes = int(resultat.group(1)), int(resultat.group(2) or 0)
        if heures <= 23 and minutes <= 59:
            return datetime.time(heures, minutes)
    # Fraction de journée Excel (cellule numérique non formatée en heure). Ex : 0.5625 = 13:30
    resultat = re.match(r"^0?\.\d+$", texte)
    if resultat:
        minutes = int(round(float(texte) * 24 * 60))
        if 0 <= minutes < 24 * 60:
            return datetime.time(minutes // 60, minutes % 60)
    raise ValueError("%s '%s' n'est pas valide (format attendu : 13:30 ou 13h30)" % (nom, texte))


def Lire_nombre(texte, nom="Le montant", entier=False, maximum=None):
    texte = texte.strip().replace(" ", "").replace(" ", "").replace("€", "").replace(",", ".")
    if not texte:
        return None
    try:
        nombre = decimal.Decimal(texte)
    except decimal.InvalidOperation:
        raise ValueError("%s '%s' n'est pas un nombre valide" % (nom, texte))
    if nombre < 0:
        raise ValueError("%s ne peut pas être négatif" % nom)
    if maximum is not None and nombre > maximum:
        raise ValueError("%s '%s' dépasse le maximum autorisé (%s)" % (nom, texte, "{:,}".format(maximum).replace(",", " ").replace(".", ",")))
    if entier:
        if nombre != nombre.to_integral_value():
            raise ValueError("%s '%s' doit être un nombre entier" % (nom, texte))
        return int(nombre)
    return nombre.quantize(decimal.Decimal("0.01"))


def Lire_booleen(texte):
    valeur = Normaliser(texte)
    if not valeur:
        return None
    if valeur in ("oui", "o", "yes", "y", "1", "x", "vrai", "true"):
        return True
    if valeur in ("non", "n", "no", "0", "faux", "false"):
        return False
    raise ValueError("La valeur '%s' n'est pas valide pour Visible sur le portail (Oui ou Non attendu)" % texte)


def Format_heure(heure):
    return heure.strftime("%Hh%M") if heure else ""


def Format_montant(montant):
    return ("%.2f %s" % (montant, utils_preferences.Get_symbole_monnaie())).replace(".", ",")


# ------------------------------------------- Lecture du fichier -------------------------------------------

def Lire_fichier(fichier=None):
    """ Lit toutes les feuilles du fichier XLSX. La première ligne non vide de chaque feuille contient les titres des colonnes """
    from openpyxl import load_workbook
    if not fichier:
        raise Erreur("Sélectionnez un fichier XLSX")
    if not fichier.name.lower().endswith((".xlsx", ".xlsm")):
        raise Erreur("Le fichier doit être au format XLSX (Excel 2007 ou ultérieur). Pour un fichier .xls ou .ods, enregistrez-le au préalable au format .xlsx.")
    if fichier.size > MAX_TAILLE_FICHIER:
        raise Erreur("Le fichier est trop volumineux (%d Mo maximum)" % (MAX_TAILLE_FICHIER // (1024 * 1024)))
    try:
        classeur = load_workbook(filename=io.BytesIO(fichier.read()), data_only=True, read_only=True)
    except Exception as err:
        logger.debug("Import événements : fichier illisible (%s)", err)
        raise Erreur("Le fichier n'a pas pu être lu. Vérifiez qu'il s'agit bien d'un fichier XLSX valide.")

    feuilles = []
    for feuille in classeur.worksheets:
        colonnes, lignes = None, []
        for num_ligne, valeurs in enumerate(feuille.iter_rows(values_only=True), start=1):
            valeurs = [Convertir_cellule(valeur) for valeur in valeurs]
            if not any(valeurs):
                continue
            if colonnes is None:
                colonnes = valeurs
                continue
            if len(lignes) >= MAX_LIGNES:
                classeur.close()
                raise Erreur("La feuille '%s' contient plus de %d lignes. Découpez le fichier en plusieurs fichiers." % (feuille.title, MAX_LIGNES))
            lignes.append({"num": num_ligne, "valeurs": valeurs})
        if colonnes is None:
            continue

        # Suppression des colonnes entièrement vides en fin de tableau et nommage des colonnes sans titre
        nbre_colonnes = max([len(colonnes)] + [len(ligne["valeurs"]) for ligne in lignes])
        while nbre_colonnes and not (nbre_colonnes <= len(colonnes) and colonnes[nbre_colonnes - 1]) and \
                not any(len(ligne["valeurs"]) >= nbre_colonnes and ligne["valeurs"][nbre_colonnes - 1] for ligne in lignes):
            nbre_colonnes -= 1
        titres, deja = [], {}
        for index in range(nbre_colonnes):
            titre = colonnes[index] if index < len(colonnes) and colonnes[index] else "Colonne %d" % (index + 1)
            deja[titre] = deja.get(titre, 0) + 1
            titres.append(titre if deja[titre] == 1 else "%s (%d)" % (titre, deja[titre]))
        for ligne in lignes:
            ligne["valeurs"] = (ligne["valeurs"] + [""] * nbre_colonnes)[:nbre_colonnes]
        feuilles.append({"nom": feuille.title, "colonnes": titres, "lignes": lignes})
    classeur.close()

    if not feuilles or not any(feuille["lignes"] for feuille in feuilles):
        raise Erreur("Le fichier ne contient aucune donnée. La première ligne doit contenir les titres des colonnes et les lignes suivantes les événements.")
    return {"nom_fichier": fichier.name, "feuilles": feuilles}


def Memoriser_fichier(session=None, idactivite=None, donnees={}):
    """ Mémorise le fichier lu en session sous un jeton propre à l'onglet (plusieurs imports peuvent être ouverts en parallèle) """
    fichiers = dict(session.get(CLE_SESSION, {}).get("fichiers", {}))
    jeton = uuid.uuid4().hex
    fichiers[jeton] = {"idactivite": idactivite, "nom_fichier": donnees["nom_fichier"], "feuilles": donnees["feuilles"]}
    while len(fichiers) > MAX_FICHIERS_SESSION:
        del fichiers[next(iter(fichiers))]
    session[CLE_SESSION] = {"fichiers": fichiers}
    return jeton


def Get_fichier_session(session=None, jeton="", idactivite=None):
    fichier = session.get(CLE_SESSION, {}).get("fichiers", {}).get(str(jeton or ""))
    if not fichier or fichier.get("idactivite") != idactivite:
        raise Erreur("Le fichier n'est plus disponible. Veuillez le sélectionner à nouveau.")
    return fichier


def Oublier_fichier(session=None, jeton=""):
    fichiers = dict(session.get(CLE_SESSION, {}).get("fichiers", {}))
    fichiers.pop(str(jeton or ""), None)
    session[CLE_SESSION] = {"fichiers": fichiers}


def Proposer_associations(colonnes=[]):
    """ Propose un champ pour chaque colonne d'après son titre. Renvoie {index_colonne: code_champ} """
    associations, deja = {}, set()
    for index, titre in enumerate(colonnes):
        titre_normalise = Normaliser(titre)
        for champ in CHAMPS:
            if champ["code"] not in deja and (titre_normalise == Normaliser(champ["label"]) or titre_normalise in champ["synonymes"]):
                associations[index] = champ["code"]
                deja.add(champ["code"])
                break
    return associations


def Get_infos_feuille(feuille={}):
    """ Données de la feuille nécessaires à l'étape d'association des colonnes """
    exemples = []
    for index in range(len(feuille["colonnes"])):
        exemple = next((ligne["valeurs"][index] for ligne in feuille["lignes"] if ligne["valeurs"][index]), "")
        if re.match(r"^\d{4}-\d{2}-\d{2}$", exemple):
            exemple = utils_dates.ConvertDateToFR(datetime.date.fromisoformat(exemple))
        exemples.append(exemple[:60])
    return {"nom": feuille["nom"], "colonnes": feuille["colonnes"], "exemples": exemples, "nbre_lignes": len(feuille["lignes"]),
            "associations": {str(index): code for index, code in Proposer_associations(feuille["colonnes"]).items()}}


# ------------------------------------------- Analyse -------------------------------------------

class Referentiel():
    """ Paramétrage de l'activité utilisé pour rapprocher les valeurs du fichier """

    def __init__(self, activite=None):
        self.activite = activite
        self.unites = list(Unite.objects.filter(activite=activite, type="Evenement").order_by("ordre"))
        self.groupes = list(Groupe.objects.filter(activite=activite).order_by("ordre"))
        self.categories = list(CategorieEvenement.objects.filter(activite=activite).order_by("nom"))
        self.dict_unites = self.Indexer(self.unites)
        self.dict_groupes = self.Indexer(self.groupes)
        self.dict_categories = self.Indexer(self.categories, abrege=False)
        self._modeles = {}

    def Indexer(self, objets=[], abrege=True):
        dictionnaire = {}
        for objet in objets:
            for valeur in (objet.nom, getattr(objet, "abrege", None) if abrege else None):
                if valeur and Normaliser(valeur) not in dictionnaire:
                    dictionnaire[Normaliser(valeur)] = objet
        return dictionnaire

    def Get_modele(self, idevenement):
        """ Evénement dont la tarification est à recopier, avec le nombre de tarifs avancés """
        if idevenement not in self._modeles:
            modele = Evenement.objects.filter(pk=idevenement, activite=self.activite).first()
            if modele:
                modele.nbre_tarifs = Tarif.objects.filter(evenement=modele).count()
            self._modeles[idevenement] = modele
        return self._modeles[idevenement]


def Lire_parametres(activite=None, parametres={}, referentiel=None):
    """ Contrôle les valeurs par défaut et les associations envoyées par la page """
    referentiel = referentiel or Referentiel(activite)
    if not referentiel.unites:
        raise Erreur("L'activité ne possède aucune unité de consommation de type Evénementiel. Créez-en une avant d'importer des événements.")
    if not referentiel.groupes:
        raise Erreur("L'activité ne possède aucun groupe. Créez-en un avant d'importer des événements.")

    unite_defaut = None
    if parametres.get("idunite"):
        unite_defaut = next((unite for unite in referentiel.unites if str(unite.pk) == str(parametres["idunite"])), None)
        if not unite_defaut:
            raise Erreur("L'unité par défaut sélectionnée n'est pas valide")
    groupes_defaut = [groupe for groupe in referentiel.groupes if str(groupe.pk) in [str(x) for x in parametres.get("idgroupes", [])]]

    associations = {}
    for index, code in (parametres.get("associations") or {}).items():
        if not code:
            continue
        if code not in DICT_CHAMPS:
            raise Erreur("Champ inconnu : %s" % code)
        if code in associations.values():
            raise Erreur("Le champ « %s » est associé à plusieurs colonnes" % DICT_CHAMPS[code]["label"])
        associations[int(index)] = code
    for code in ("date", "nom"):
        if code not in associations.values():
            raise Erreur("Associez une colonne au champ obligatoire « %s »" % DICT_CHAMPS[code]["label"])
    if not unite_defaut and "unite" not in associations.values():
        raise Erreur("Sélectionnez une unité par défaut ou associez une colonne au champ « Unité »")
    if not groupes_defaut and "groupe" not in associations.values():
        raise Erreur("Sélectionnez au moins un groupe par défaut ou associez une colonne au champ « Groupe(s) »")
    return referentiel, unite_defaut, groupes_defaut, associations


def Analyser(activite=None, feuille={}, parametres={}):
    """ Contrôle chaque ligne de la feuille et prépare les événements à créer """
    referentiel, unite_defaut, groupes_defaut, associations = Lire_parametres(activite, parametres)
    index_champs = {code: index for index, code in associations.items()}

    def Valeur(ligne, code):
        index = index_champs.get(code)
        return ligne["valeurs"][index].strip() if index is not None and index < len(ligne["valeurs"]) else ""

    # Lecture de chaque ligne
    lignes_lues = []
    for ligne in feuille["lignes"]:
        resultat = {"num": ligne["num"], "erreurs": [], "avertissements": [], "valeurs": {}, "groupes": []}
        valeurs = resultat["valeurs"]
        try:
            valeurs["date"] = Lire_date(Valeur(ligne, "date"))
            if not valeurs["date"]:
                resultat["erreurs"].append("La date est obligatoire")
        except ValueError as err:
            resultat["erreurs"].append(str(err))

        valeurs["nom"] = Valeur(ligne, "nom")
        if not valeurs["nom"]:
            resultat["erreurs"].append("Le nom est obligatoire")
        elif len(valeurs["nom"]) > 200:
            resultat["erreurs"].append("Le nom dépasse 200 caractères")

        valeurs["description"] = Valeur(ligne, "description") or None
        if valeurs["description"] and len(valeurs["description"]) > 400:
            resultat["erreurs"].append("La description dépasse 400 caractères")

        # Unité
        texte = Valeur(ligne, "unite")
        valeurs["unite"] = referentiel.dict_unites.get(Normaliser(texte)) if texte else unite_defaut
        if texte and not valeurs["unite"]:
            resultat["erreurs"].append("L'unité '%s' n'existe pas ou n'est pas de type Evénementiel" % texte)
        elif not valeurs["unite"]:
            resultat["erreurs"].append("Aucune unité renseignée et aucune unité par défaut")

        # Groupes
        texte = Valeur(ligne, "groupe")
        if texte:
            if Normaliser(texte) in ("tous", "tous les groupes"):
                resultat["groupes"] = list(referentiel.groupes)
            else:
                for nom_groupe in [x.strip() for x in texte.split(";") if x.strip()]:
                    groupe = referentiel.dict_groupes.get(Normaliser(nom_groupe))
                    if not groupe:
                        resultat["erreurs"].append("Le groupe '%s' n'existe pas dans l'activité" % nom_groupe)
                    elif groupe not in resultat["groupes"]:
                        resultat["groupes"].append(groupe)
        else:
            resultat["groupes"] = list(groupes_defaut)
        if not resultat["groupes"] and not any("groupe" in erreur for erreur in resultat["erreurs"]):
            resultat["erreurs"].append("Aucun groupe renseigné et aucun groupe par défaut")

        # Catégorie
        texte = Valeur(ligne, "categorie")
        valeurs["categorie"] = referentiel.dict_categories.get(Normaliser(texte)) if texte else None
        if texte and not valeurs["categorie"]:
            resultat["erreurs"].append("La catégorie '%s' n'existe pas dans l'activité" % texte)

        # Heures
        for code, nom in (("heure_debut", "L'heure de début"), ("heure_fin", "L'heure de fin"), ("equiv_heures", "L'équivalence en heures")):
            try:
                valeurs[code] = Lire_heure(Valeur(ligne, code), nom=nom)
            except ValueError as err:
                valeurs[code] = None
                resultat["erreurs"].append(str(err))
        if valeurs["heure_debut"] and valeurs["heure_fin"] and valeurs["heure_fin"] <= valeurs["heure_debut"]:
            resultat["erreurs"].append("L'heure de fin doit être postérieure à l'heure de début")

        # Capacité et montant
        try:
            valeurs["capacite_max"] = Lire_nombre(Valeur(ligne, "capacite_max"), nom="La capacité", entier=True, maximum=MAX_CAPACITE) or None
        except ValueError as err:
            valeurs["capacite_max"] = None
            resultat["erreurs"].append(str(err))
        try:
            valeurs["montant"] = Lire_nombre(Valeur(ligne, "montant"), nom="Le montant", maximum=MAX_MONTANT)
        except ValueError as err:
            valeurs["montant"] = None
            resultat["erreurs"].append(str(err))

        # Tarification recopiée d'un événement existant (prioritaire sur le montant)
        valeurs["modele"] = None
        texte = Valeur(ligne, "tarif_existant")
        if texte:
            try:
                idmodele = Lire_nombre(texte, nom="L'ID de l'événement modèle", entier=True)
            except ValueError as err:
                idmodele = None
                resultat["erreurs"].append(str(err))
            if idmodele is not None:
                modele = referentiel.Get_modele(idmodele) if idmodele <= MAX_ID else None
                if not modele:
                    resultat["erreurs"].append("L'événement modèle ID%d n'existe pas dans cette activité" % idmodele)
                else:
                    valeurs["modele"] = modele
                    if valeurs["montant"]:
                        resultat["avertissements"].append("Montant du fichier ignoré : tarification recopiée depuis l'événement ID%d" % idmodele)
                    valeurs["montant"] = modele.montant or None

        # Visibilité sur le portail
        try:
            visible = Lire_booleen(Valeur(ligne, "visible_portail"))
            valeurs["visible_portail"] = True if visible is None else visible
        except ValueError as err:
            valeurs["visible_portail"] = True
            resultat["erreurs"].append(str(err))

        lignes_lues.append(resultat)

    # Evénements et ouvertures déjà existants sur la période du fichier
    dates = [ligne["valeurs"]["date"] for ligne in lignes_lues if ligne["valeurs"].get("date")]
    existants, ouvertures = set(), set()
    if dates:
        for date, nom, idunite, idgroupe in Evenement.objects.filter(activite=activite, date__range=(min(dates), max(dates))).values_list("date", "nom", "unite_id", "groupe_id"):
            existants.add((date, Normaliser(nom), idunite, idgroupe))
        ouvertures = set(Ouverture.objects.filter(activite=activite, date__range=(min(dates), max(dates))).values_list("date", "unite_id", "groupe_id"))

    # Détail des événements à créer (un par groupe)
    resultats, a_creer, cles_fichier, ouvertures_a_creer = [], [], set(), set()
    compteurs = {"lignes": len(lignes_lues), "a_creer": 0, "doublons": 0, "erreurs": 0, "ouvertures": 0, "tarifs_avances": 0}
    for ligne in lignes_lues:
        valeurs = ligne["valeurs"]
        base = {
            "ligne": ligne["num"], "date": utils_dates.ConvertDateToFR(valeurs["date"]) if valeurs.get("date") else "", "nom": valeurs.get("nom", ""),
            "unite": valeurs["unite"].nom if valeurs.get("unite") else "",
            "horaires": "-".join([x for x in (Format_heure(valeurs.get("heure_debut")), Format_heure(valeurs.get("heure_fin"))) if x]),
            "tarification": Get_texte_tarification(valeurs),
            "avertissements": ligne["avertissements"],
        }
        if ligne["erreurs"]:
            compteurs["erreurs"] += 1
            resultats.append(dict(base, groupe=", ".join(groupe.nom for groupe in ligne["groupes"]), statut="erreur", message=" ; ".join(ligne["erreurs"])))
            continue
        for groupe in ligne["groupes"]:
            cle = (valeurs["date"], Normaliser(valeurs["nom"]), valeurs["unite"].pk, groupe.pk)
            detail = dict(base, groupe=groupe.nom)
            if cle in existants:
                compteurs["doublons"] += 1
                resultats.append(dict(detail, statut="doublon", message="Existe déjà, ignoré"))
                continue
            if cle in cles_fichier:
                compteurs["doublons"] += 1
                resultats.append(dict(detail, statut="doublon", message="Doublon dans le fichier, ignoré"))
                continue
            cles_fichier.add(cle)
            cle_ouverture = (valeurs["date"], valeurs["unite"].pk, groupe.pk)
            ouverture = cle_ouverture not in ouvertures and cle_ouverture not in ouvertures_a_creer
            if ouverture:
                ouvertures_a_creer.add(cle_ouverture)
            compteurs["a_creer"] += 1
            if valeurs["modele"] and valeurs["modele"].nbre_tarifs and not valeurs["modele"].montant:
                compteurs["tarifs_avances"] += 1
            resultats.append(dict(detail, statut="creer", message="À créer" + (" + ouverture du calendrier" if ouverture else "")))
            a_creer.append({"valeurs": valeurs, "groupe": groupe})
    compteurs["ouvertures"] = len(ouvertures_a_creer)
    return {"compteurs": compteurs, "lignes": resultats, "a_creer": a_creer, "ouvertures": ouvertures_a_creer}


def Get_texte_tarification(valeurs={}):
    modele = valeurs.get("modele")
    if modele:
        if modele.montant:
            return "%s (copie ID%d)" % (Format_montant(modele.montant), modele.pk)
        if modele.nbre_tarifs:
            pluriel = "s" if modele.nbre_tarifs > 1 else ""
            return "%d tarif%s avancé%s (copie ID%d)" % (modele.nbre_tarifs, pluriel, pluriel, modele.pk)
        return "Gratuit (copie ID%d)" % modele.pk
    if valeurs.get("montant"):
        return Format_montant(valeurs["montant"])
    return "Gratuit"


def Get_apercu(activite=None, feuille={}, parametres={}):
    """ Résultat de l'analyse renvoyé à la page (sans les objets) """
    analyse = Analyser(activite, feuille, parametres)
    return {"compteurs": analyse["compteurs"], "lignes": analyse["lignes"]}


# ------------------------------------------- Importation -------------------------------------------

def Copier_tarifs(evenement=None, modele=None, tarifs_modele=None):
    """ Duplique les tarifs avancés de l'événement modèle (même logique que la saisie d'un événement).
    Si un tarif est réservé au groupe de l'événement modèle, il est réservé au groupe du nouvel événement :
    sinon, un événement créé dans un autre groupe recevrait un tarif qui ne s'appliquerait jamais. """
    for tarif in tarifs_modele:
        nouveau_tarif = Tarif.objects.get(pk=tarif.pk)
        nouveau_tarif.pk = None
        nouveau_tarif.evenement = evenement
        nouveau_tarif.date_debut = evenement.date
        nouveau_tarif.date_fin = evenement.date
        nouveau_tarif.save()
        for champ in Tarif._meta.many_to_many:
            valeurs = list(getattr(tarif, champ.name).all())
            if champ.name == "groupes" and evenement.groupe_id != modele.groupe_id and modele.groupe_id in [groupe.pk for groupe in valeurs]:
                valeurs = [groupe for groupe in valeurs if groupe.pk != modele.groupe_id] + [evenement.groupe]
            getattr(nouveau_tarif, champ.name).add(*valeurs)
        for ligne_tarif in TarifLigne.objects.filter(tarif=tarif):
            ligne_tarif.pk = None
            ligne_tarif.tarif = nouveau_tarif
            ligne_tarif.save()


def Importer(activite=None, feuille={}, parametres={}, request=None, nom_fichier=""):
    tarifs_modeles = {}
    with transaction.atomic():
        # Verrou sur l'activité : deux importations simultanées sont traitées l'une après l'autre,
        # la seconde voit alors les événements de la première comme des doublons
        Activite.objects.select_for_update().filter(pk=activite.pk).first()
        analyse = Analyser(activite, feuille, parametres)
        if not analyse["a_creer"]:
            raise Erreur("Aucun événement à créer")

        # Ouvertures manquantes dans le calendrier
        Ouverture.objects.bulk_create([Ouverture(activite=activite, date=date, unite_id=idunite, groupe_id=idgroupe) for date, idunite, idgroupe in analyse["ouvertures"]])

        # Evénements
        for element in analyse["a_creer"]:
            valeurs, modele = element["valeurs"], element["valeurs"]["modele"]
            evenement = Evenement.objects.create(
                activite=activite, unite=valeurs["unite"], groupe=element["groupe"], date=valeurs["date"], nom=valeurs["nom"],
                description=valeurs["description"], capacite_max=valeurs["capacite_max"], heure_debut=valeurs["heure_debut"],
                heure_fin=valeurs["heure_fin"], montant=valeurs["montant"] or None, equiv_heures=valeurs["equiv_heures"],
                visible_portail=valeurs["visible_portail"], categorie=valeurs["categorie"],
            )
            # Recopie des tarifs avancés de l'événement modèle
            if modele and not modele.montant and modele.nbre_tarifs:
                if modele.pk not in tarifs_modeles:
                    tarifs_modeles[modele.pk] = list(Tarif.objects.filter(evenement=modele).prefetch_related(*[champ.name for champ in Tarif._meta.many_to_many]))
                Copier_tarifs(evenement=evenement, modele=modele, tarifs_modele=tarifs_modeles[modele.pk])

        # Historique
        compteurs = analyse["compteurs"]
        detail = "Fichier=%s, Evénements créés=%d, Ouvertures créées=%d, Doublons ignorés=%d, Lignes en erreur=%d" % (
            nom_fichier, compteurs["a_creer"], compteurs["ouvertures"], compteurs["doublons"], compteurs["erreurs"])
        utils_historique.Ajouter(titre="Importation d'événements", detail=detail, utilisateur=request.user if request else None, activite=activite.pk)

    logger.debug("Import événements : %d événements créés dans l'activité %d", compteurs["a_creer"], activite.pk)
    return {"compteurs": compteurs, "erreurs": [ligne for ligne in analyse["lignes"] if ligne["statut"] == "erreur"]}


# ------------------------------------------- Fichier exemple -------------------------------------------

def Get_fichier_exemple(activite=None):
    """ Fichier XLSX modèle avec les titres de colonnes reconnus automatiquement et des exemples tirés du paramétrage """
    import xlsxwriter
    referentiel = Referentiel(activite)
    unite = referentiel.unites[0].nom if referentiel.unites else "Sortie"
    groupes = [groupe.nom for groupe in referentiel.groupes] or ["Groupe unique"]
    annee = datetime.date.today().year + 1

    sortie = io.BytesIO()
    classeur = xlsxwriter.Workbook(sortie, {"in_memory": True})
    titre = classeur.add_format({"bold": True, "bg_color": "#DDEBF7", "border": 1})
    format_date = classeur.add_format({"num_format": "dd/mm/yyyy"})
    format_heure = classeur.add_format({"num_format": "hh:mm"})
    format_montant = classeur.add_format({"num_format": "0.00"})

    feuille = classeur.add_worksheet("Evénements")
    codes = ["date", "nom", "description", "unite", "groupe", "heure_debut", "heure_fin", "capacite_max", "montant", "tarif_existant", "visible_portail"]
    for colonne, code in enumerate(codes):
        feuille.write(0, colonne, DICT_CHAMPS[code]["label"], titre)
    exemples = [
        [datetime.date(annee, 2, 12), "Patinoire", "RDV à 13h15 devant le centre. Prévoir des gants.", unite, groupes[-1], datetime.time(13, 30), datetime.time(17, 0), 24, 6.5, None, "Oui"],
        [datetime.date(annee, 2, 13), "Cinéma", None, unite, ";".join(groupes[:2]), datetime.time(14, 0), datetime.time(16, 0), None, 4, None, "Oui"],
        [datetime.date(annee, 2, 14), "Grand jeu", "Prévoir une tenue de sport.", None, None, None, None, None, None, None, "Non"],
    ]
    for ligne, valeurs in enumerate(exemples, start=1):
        for colonne, valeur in enumerate(valeurs):
            if isinstance(valeur, datetime.date):
                feuille.write_datetime(ligne, colonne, datetime.datetime.combine(valeur, datetime.time()), format_date)
            elif isinstance(valeur, datetime.time):
                feuille.write_datetime(ligne, colonne, valeur, format_heure)
            elif isinstance(valeur, float):
                feuille.write_number(ligne, colonne, valeur, format_montant)
            elif valeur is not None:
                feuille.write(ligne, colonne, valeur)
    feuille.set_column(0, 0, 12)
    feuille.set_column(1, 1, 24)
    feuille.set_column(2, 2, 45)
    feuille.set_column(3, 4, 22)
    feuille.set_column(5, 8, 14)
    feuille.set_column(9, 9, 38)
    feuille.set_column(10, 10, 22)
    feuille.freeze_panes(1, 0)

    aide = classeur.add_worksheet("Aide")
    aide.write_row(0, 0, ["Colonne", "Obligatoire", "Contenu attendu"], titre)
    for ligne, champ in enumerate(CHAMPS, start=1):
        aide.write_row(ligne, 0, [champ["label"], "Oui" if champ.get("obligatoire") else "Non", champ["aide"]])
    aide.set_column(0, 0, 38)
    aide.set_column(1, 1, 12)
    aide.set_column(2, 2, 120)
    classeur.close()
    return sortie.getvalue()
